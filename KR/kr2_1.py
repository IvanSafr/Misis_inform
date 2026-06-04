import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side


def create_lab_excel():
    # Создаем новую рабочую книгу
    wb = openpyxl.Workbook()

    # Стили для оформления
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    def apply_style(ws, min_row, max_row, min_col, max_col):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align

    # ==========================================
    # ЛИСТ 1: ТАБЛИЦА 1 (Определение модуля кручения)
    # ==========================================
    ws1 = wb.active
    ws1.title = "Таблица 1"

    # Вводные погрешности и параметры (из текста)
    ws1['A1'] = "Параметры установки и абсолютные погрешности (п. 5.2):"
    ws1['A1'].font = bold_font
    ws1.append(["h (м) - плечо силы", 0.15])  # h принимаем равным радиусу диска
    ws1.append(["ΔF (Н)", 0.1])
    ws1.append(["Δh (м)", 0.001])
    ws1.append(["Δφ (рад)", 0.0436])
    ws1.append([])  # Пустая строка

    # Заголовки таблицы 1
    headers1 = ["№ п/п", "F_i, Н", "φ_i, рад", "M_i = F_i*h, Н*м", "N_i = M_i/φ_i, Н*м/рад", "ΔN_i = |N_i - N_ср|"]
    ws1.append(headers1)
    for col in range(1, len(headers1) + 1):
        ws1.cell(row=7, column=col).font = bold_font
        ws1.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

    # Строки для измерений (12 опытов)
    start_row = 8
    end_row = 19
    for i in range(1, 13):
        row_num = start_row + i - 1
        ws1.cell(row=row_num, column=1, value=i)
        # Данные для F_i и phi_i вводятся вручную (оставляем пустыми)
        # Формулы:
        ws1.cell(row=row_num, column=4, value=f"=B{row_num}*$B$2")  # M_i
        ws1.cell(row=row_num, column=5, value=f"=IF(C{row_num}=0,\"\",D{row_num}/C{row_num})")  # N_i
        ws1.cell(row=row_num, column=6, value=f"=IF(E{row_num}=\"\",\"\",ABS(E{row_num}-$E$20))")  # ΔN_i

    apply_style(ws1, 7, end_row, 1, 6)

    # Итоговые расчеты под таблицей 1
    ws1['D20'] = "СРЕДНЕЕ (N_ср):"
    ws1['E20'] = f"=AVERAGE(E{start_row}:E{end_row})"

    ws1['D21'] = "Относ. погр. δ_N:"
    # Формула из методички: δ_N = ΔF/F_ср + Δh/h + Δφ/φ_ср
    ws1['E21'] = f"=$B$3/AVERAGE(B{start_row}:B{end_row}) + $B$4/$B$2 + $B$5/AVERAGE(C{start_row}:C{end_row})"

    ws1['D22'] = "Абс. погр. ΔN_ср:"
    ws1['E22'] = "=E20*E21"  # ΔN = N_ср * δ_N

    # ==========================================
    # ЛИСТ 2: ТАБЛИЦА 2 (Основной опыт)
    # ==========================================
    ws2 = wb.create_sheet(title="Таблица 2")

    headers2 = ["№ п/п", "a1, м", "a2, м", "T_i, с", "T_ср, с", "ΔT_i, с", "ΔT_ср, с", "δ_T, %"]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        ws2.cell(row=1, column=col).font = bold_font
        ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    # Добавим 6 строк (по 3 измерения на 2 варианта установки дополнительных цилиндров)
    for i in range(1, 7):
        ws2.cell(row=i + 1, column=1, value=i)
        ws2.cell(row=i + 1, column=6, value=f"=IF(D{i + 1}=\"\",\"\",ABS(D{i + 1}-E{i + 1}))")  # ΔT_i

    # Группируем расчет средних для опытов по 3 замера
    ws2['E2'] = "=AVERAGE(D2:D4)"  # T_ср для первой серии
    ws2['E3'] = "=E2"
    ws2['E4'] = "=E2"
    ws2['G2'] = "=AVERAGE(F2:F4)"  # ΔT_ср для первой серии
    ws2['G3'] = "=G2"
    ws2['G4'] = "=G2"
    ws2['H2'] = "=(G2/E2)*100"  # δ_T в процентах

    ws2['E5'] = "=AVERAGE(D5:D7)"  # T_ср для второй серии
    ws2['E6'] = "=E5"
    ws2['E7'] = "=E5"
    ws2['G5'] = "=AVERAGE(F5:F7)"
    ws2['G6'] = "=G5"
    ws2['G7'] = "=G5"
    ws2['H5'] = "=(G5/E5)*100"

    apply_style(ws2, 1, 7, 1, 8)

    # ==========================================
    # ЛИСТ 3: ТАБЛИЦА 3 (Моменты инерции. Теорема Штейнера)
    # ==========================================
    ws3 = wb.create_sheet(title="Таблицы 3.1 и 3.2")
    ws3.column_dimensions['A'].width = 25
    ws3.column_dimensions['B'].width = 15
    ws3.column_dimensions['C'].width = 15
    ws3.column_dimensions['D'].width = 20
    ws3.column_dimensions['E'].width = 20
    ws3.column_dimensions['F'].width = 20
    ws3.column_dimensions['G'].width = 15
    ws3.column_dimensions['H'].width = 15

    # Дано из примечания (массы переведены в кг, радиусы в метры)
    ws3['A1'] = "Параметры тел (из методички):"
    ws3['A1'].font = bold_font
    ws3.append(["Тело", "Масса m, кг", "Радиус r, м"])
    ws3.append(["Основной диск", 0.425, 0.150])
    ws3.append(["Цилиндр №1", 0.198, 0.0199])
    ws3.append(["Цилиндр №2", 0.160, 0.01505])
    apply_style(ws3, 2, 5, 1, 3)
    ws3.append([])

    # Таблица 3 (Теоретический и Экспериментальный расчеты)
    ws3.append(["Сравнение результатов (для Серии 1 из Табл. 2)"])
    ws3.cell(row=7, column=1).font = bold_font

    headers3 = ["Тело, входящее в систему", "a, м (расст.)", "I_0, кг*м2", "I_теор, кг*м2", "I_эксп, кг*м2",
                "Абс. погр. ΔI_э", "Отн. погр. δ_I, %"]
    ws3.append(headers3)
    for col in range(1, len(headers3) + 1):
        ws3.cell(row=8, column=col).font = bold_font

    # Диск (a=0)
    ws3.cell(row=9, column=1, value="Диск")
    ws3.cell(row=9, column=2, value=0)
    ws3.cell(row=9, column=3, value="=0.5*B3*(C3^2)")  # I_0 = 1/2 * m * r^2
    ws3.cell(row=9, column=4, value="=C9")  # I_теор = I_0 + m*0^2

    # Цилиндр №1 (a берется из первой строки Табл 2)
    ws3.cell(row=10, column=1, value="Цилиндр №1")
    ws3.cell(row=10, column=2, value="='Таблица 2'!B2")
    ws3.cell(row=10, column=3, value="=0.5*B4*(C4^2)")  # I_0
    ws3.cell(row=10, column=4, value="=C10 + B4*(B10^2)")  # I_теор = I_0 + m*a^2 (Теорема Штейнера)

    # Цилиндр №2
    ws3.cell(row=11, column=1, value="Цилиндр №2")
    ws3.cell(row=11, column=2, value="='Таблица 2'!C2")
    ws3.cell(row=11, column=3, value="=0.5*B5*(C5^2)")  # I_0
    ws3.cell(row=11, column=4, value="=C11 + B5*(B11^2)")  # I_теор = I_0 + m*a^2

    # Система (Сумма)
    ws3.cell(row=12, column=1, value="СИСТЕМА ТЕЛ")
    ws3.cell(row=12, column=1).font = bold_font
    ws3.cell(row=12, column=4, value="=SUM(D9:D11)")  # I_теор всей системы (сумма)

    # Экспериментальный момент инерции I_э = T_ср^2 * N_ср / (4 * pi^2)
    ws3.cell(row=12, column=5, value="=(('Таблица 2'!E2^2)/(4*PI()^2)) * 'Таблица 1'!E20")

    # Относительная погрешность δ_I = δ_N + 2*(ΔT/T)
    ws3.cell(row=12, column=7, value="='Таблица 1'!E21 + 2*('Таблица 2'!G2/'Таблица 2'!E2)")
    ws3.cell(row=12, column=7).number_format = '0.00%'

    # Абсолютная погрешность ΔI = I_э * δ_I
    ws3.cell(row=12, column=6, value="=E12*G12")

    apply_style(ws3, 8, 12, 1, 7)

    # Сохраняем файл
    file_name = "Lab_1-04_Steiner.xlsx"
    wb.save(file_name)
    print(f"Файл '{file_name}' успешно создан! Все формулы настроены по теории лабораторной работы.")


if __name__ == "__main__":
    create_lab_excel()