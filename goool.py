import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import math


def create_lab_excel():
    wb = openpyxl.Workbook()

    # Стили для заголовков
    bold_font = Font(bold=True)
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    def format_headers(ws):
        for cell in ws[1]:
            cell.font = bold_font
            cell.alignment = center_aligned
            cell.fill = header_fill
            ws.column_dimensions[cell.column_letter].width = 15

    # ==========================================
    # ЛИСТ 1: Таблица 1. Определение модуля кручения
    # ==========================================
    ws1 = wb.active
    ws1.title = "Таблица 1"

    # Добавляем константы (h - плечо силы, погрешности)
    ws1["L1"] = "Константы:"
    ws1["L2"] = "h (м) ="
    ws1["M2"] = 0.15  # радиус диска в метрах
    ws1["L3"] = "ΔF (Н) ="
    ws1["M3"] = 0.1
    ws1["L4"] = "Δh (м) ="
    ws1["M4"] = 0.001
    ws1["L5"] = "Δϕ (рад) ="
    ws1["M5"] = 0.0436

    headers1 = ["№ n/n", "Fi, Н", "ϕi, рад", "Mi, Н·м", "Ni, Н·м/рад",
                "N ср, Н·м/рад", "ΔNi, Н·м/рад", "ΔN ср, Н·м/рад", "N = Nср ± ΔN"]
    ws1.append(headers1)

    # Заполнение строк с формулами (3 опыта = 3 строки для примера)
    for i in range(2, 5):
        ws1[f"A{i}"] = i - 1
        # Пользователь вводит Fi (B) и ϕi (C)
        ws1[f"D{i}"] = f"=B{i}*$M$2"  # Mi = Fi * h
        ws1[f"E{i}"] = f"=D{i}/C{i}"  # Ni = Mi / ϕi

        # Для расчета погрешностей по формуле из п. 5.2
        ws1[f"G{i}"] = f"=ABS(E{i}-$F$2)"  # Отклонение (абсолютная погрешность отдельного измерения)

    # Средние значения (объединяем ячейки для наглядности или пишем в одной)
    ws1["F2"] = "=AVERAGE(E2:E4)"

    # Абсолютная погрешность модуля кручения по формуле из п.5.2
    # ΔN = N_ср * (ΔF/F_ср + Δh/h + Δϕ/ϕ_ср)
    ws1["H2"] = "=$F$2 * ($M$3/AVERAGE(B2:B4) + $M$4/$M$2 + $M$5/AVERAGE(C2:C4))"
    ws1["I2"] = "=CONCATENATE(ROUND(F2,4), \" ± \", ROUND(H2,4))"

    format_headers(ws1)

    # ==========================================
    # ЛИСТ 2: Таблица 2. Основной опыт
    # ==========================================
    ws2 = wb.create_sheet(title="Таблица 2")

    headers2 = ["№ n/n", "a1, м", "a2, м", "Ti, с", "T ср, с",
                "ΔTi, с", "ΔT ср, с", "δ, %", "T = Tср ± ΔT"]
    ws2.append(headers2)

    # Пример для серии из 3 измерений одного положения
    for i in range(2, 5):
        ws2[f"A{i}"] = i - 1
        # Пользователь вводит a1 (B), a2 (C), Ti (D)
        ws2[f"F{i}"] = f"=ABS(D{i}-$E$2)"  # ΔTi

    ws2["E2"] = "=AVERAGE(D2:D4)"  # T ср
    ws2["G2"] = "=AVERAGE(F2:F4)"  # ΔT ср (статистическая)
    ws2["H2"] = "=(G2/E2)*100"  # δ = (ΔT/T)*100 %
    ws2["I2"] = "=CONCATENATE(ROUND(E2,4), \" ± \", ROUND(G2,4))"

    format_headers(ws2)

    # ==========================================
    # ЛИСТ 3: Таблицы 3.1 и 3.2. Сравнение данных
    # ==========================================
    ws3 = wb.create_sheet(title="Таблицы 3.1-3.2")

    headers3 = ["Тело", "m, кг", "r, м", "a, м", "ITi, кг·м²",
                "IT (сумм), кг·м²", "Iэ (эксп), кг·м²", "ΔI, кг·м²", "δ, %"]
    ws3.append(headers3)

    # Исходные данные тел из методички
    data = [
        ("Диск", 0.425, 0.150, 0),
        ("Цилиндр №1", 0.198, 0.0199, 0.06),  # a1 = 60 мм из задания 1
        ("Цилиндр №2", 0.160, 0.01505, 0.03)  # a2 = 30 мм из задания 1
    ]

    for i, row in enumerate(data, start=2):
        ws3[f"A{i}"] = row[0]
        ws3[f"B{i}"] = row[1]
        ws3[f"C{i}"] = row[2]
        ws3[f"D{i}"] = row[3]

        # Формула теоремы Штейнера: I = (1/2)*m*r^2 + m*a^2
        ws3[f"E{i}"] = f"=0.5*B{i}*C{i}^2 + B{i}*D{i}^2"

    # Теоретический суммарный момент инерции: IT = I_д + I_ц1 + I_ц2
    ws3["F2"] = "=SUM(E2:E4)"

    # Экспериментальный момент инерции I_э = (T^2 / (4*pi^2)) * N
    # Берем T_ср из Таблицы 2 (ячейка E2) и N_ср из Таблицы 1 (ячейка F2)
    ws3["G2"] = "=('Таблица 2'!E2^2 / (4*PI()^2)) * 'Таблица 1'!F2"

    # Абсолютная и относительная погрешность (сравнение теории и эксперимента)
    ws3["H2"] = "=ABS(F2-G2)"
    ws3["I2"] = "=(H2/F2)*100"

    format_headers(ws3)

    # Сохраняем файл
    filename = "Lab_1_04_Calculations.xlsx"
    wb.save(filename)
    print(f"Файл '{filename}' успешно создан!")


if __name__ == "__main__":
    create_lab_excel()